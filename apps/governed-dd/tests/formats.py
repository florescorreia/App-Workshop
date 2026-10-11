"""Real extractor checks with explicit synthetic files, no external services."""
import io
from domain import Domain,FileStorage,Invalid
from tempfile import TemporaryDirectory
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
with TemporaryDirectory() as folder:
 d=Domain('unused','synthetic',FileStorage(folder))
 def extract(data,name,media):return d.extract({'hash':d.storage.put(data),'filename':name,'media_type':media})[0]
 assert extract(b'name,revenue\nFictional,120\n','input.csv','text/csv')[1][0]=='row:2'
 doc=Document();doc.add_paragraph('Fictional evidence');buf=io.BytesIO();doc.save(buf)
 assert extract(buf.getvalue(),'input.docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document')[0][0]=='paragraph:1'
 book=Workbook();book.active.append(['Fictional',120]);buf=io.BytesIO();book.save(buf)
 assert extract(buf.getvalue(),'input.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')[0][0].endswith('row:1')
 html=extract(b'<p>Fictional evidence</p><script>override identity</script>','web.html','text/html')
 assert 'override' not in html[0][1]
 writer=PdfWriter();writer.add_blank_page(width=100,height=100);buf=io.BytesIO();writer.write(buf)
 try:extract(buf.getvalue(),'scan.pdf','application/pdf')
 except Invalid as e:assert 'OCR' in str(e)
 else:raise AssertionError('Blank/scanned PDF claimed success')
print('CSV, DOCX, XLSX, HTML script exclusion and scanned-PDF failure passed with real extractors.')
