"""Canonical tenant service. API and MCP call these same operations."""
import csv
import hashlib
import io
import json
import os
import socket
import ipaddress
import uuid
from pathlib import Path
from urllib.parse import urlparse

import httpx
import psycopg
import yaml
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb
from neo4j import GraphDatabase


class Invalid(ValueError):
    pass


KINDS = {'Source','Evidence','Company','Person','Market','Product','Geography','Metric',
         'Definition','Finding','Assumption','AnalysisRun','Artifact','Task','Assertion','Deliverable','Review','Relation','ReviewerValidation','OntologyProposal'}
REQUIRED = {
    'Source': ['filename','hash','size','media_type','retrieved_at','access_limitations'],
    'Finding': ['claim','implications','evidence','assumptions','uncertainty','run'],
    'Definition': ['name','meaning','units','scope','owner'],
    'AnalysisRun': ['task_id','paperclip_instance','paperclip_run_id','inputs','code_commit','environment','parameters','seeds','outputs','execution_status'],
    'Artifact': ['hash','size','media_type','run'],
    'Deliverable': ['title','findings','artifacts','run'],
    'Assertion': ['type','properties','time_scope','evidence','run','confidence','review_state'],
    'Relation': ['relation','from','to','evidence','run'],
    'ReviewerValidation': ['reviewed','report','checks','reviewer_run'],
    'OntologyProposal': ['document','rationale'],
}


def canonical(body):
    return json.dumps(body, sort_keys=True, separators=(',', ':'), ensure_ascii=False).encode()


def public_url(url):
    p = urlparse(url)
    if p.scheme != 'https' or not p.hostname or p.username or p.password or p.port not in (None,443):
        raise Invalid('Only public HTTPS sources are permitted')
    try:
        addresses = socket.getaddrinfo(p.hostname,443,type=socket.SOCK_STREAM)
    except OSError as exc:
        raise Invalid('Source DNS resolution failed') from exc
    if not addresses or any(not ipaddress.ip_address(a[4][0]).is_global for a in addresses):
        raise Invalid('Source destination is not public')
    return p.hostname, sorted({a[4][0] for a in addresses})


def fetch_public(url, limit=20*1024*1024):
    """Pin DNS with a direct socket connection; preserve HTTPS hostname verification.

    No proxy env, DNS rebinding window or redirect to an unchecked destination.
    """
    import ssl
    import http.client
    for _ in range(4):
        host, addresses = public_url(url)
        p = urlparse(url)
        class Pinned(http.client.HTTPSConnection):
            def connect(self):
                raw = socket.create_connection((addresses[0],443),timeout=20)
                context=ssl.create_default_context()
                if os.getenv('PUBLIC_FETCH_CA_FILE'):context.load_verify_locations(os.environ['PUBLIC_FETCH_CA_FILE'])
                self.sock = context.wrap_socket(raw,server_hostname=host)
        conn = Pinned(host,timeout=20)
        try:
            conn.request('GET',p.path + ('?'+p.query if p.query else ''),headers={'User-Agent':'GovernedDD/0.1 source capture','Accept-Encoding':'identity'})
            response = conn.getresponse()
            if response.status in (301,302,303,307,308):
                from urllib.parse import urljoin
                url = urljoin(url,response.getheader('location',''))
                continue
            if response.status != 200:
                raise Invalid(f'Source unavailable (HTTP {response.status})')
            data = response.read(limit+1)
            if len(data)>limit:
                raise Invalid('Source exceeds download limit')
            return data,response.getheader('Content-Type','application/octet-stream').split(';')[0],url
        finally:
            conn.close()
    raise Invalid('Too many redirects')


class FileStorage:
    def __init__(self,root='/artifacts'):
        self.root = Path(root)
        self.root.mkdir(parents=True,exist_ok=True)

    def put(self,data):
        digest = hashlib.sha256(data).hexdigest()
        path = self.root/digest[:2]/digest
        path.parent.mkdir(exist_ok=True)
        try:
            with path.open('xb') as output:
                output.write(data)
            path.chmod(0o440)
        except FileExistsError:
            if hashlib.sha256(path.read_bytes()).hexdigest()!=digest:
                raise Invalid('Stored bytes failed checksum')
        return digest

    def get(self,digest):
        if len(digest)!=64 or any(c not in '0123456789abcdef' for c in digest):
            raise Invalid('Invalid content hash')
        try:
            data=(self.root/digest[:2]/digest).read_bytes()
        except FileNotFoundError as exc:
            raise Invalid('Artifact bytes unavailable') from exc
        if hashlib.sha256(data).hexdigest()!=digest:
            raise Invalid('Artifact checksum failed')
        return data


class Domain:
    def __init__(self,dsn=None,tenant=None,storage=None):
        if os.getenv('PROVIDER_KEYS_FILE'):
            allowed={os.getenv('WEB_SEARCH_API_KEY_ENV','DD_SEARCH_API_KEY'),os.getenv('EMBEDDING_API_KEY_ENV','DD_EMBEDDING_API_KEY')}
            for name,value in json.loads(Path(os.environ['PROVIDER_KEYS_FILE']).read_text()).items():
                if name in allowed:os.environ[name]=value
        self.dsn = dsn or os.environ['DATABASE_URL']
        self.tenant = tenant or os.environ['TENANT_ID']
        self.storage = storage or FileStorage()

    def connect(self):
        return psycopg.connect(self.dsn,row_factory=dict_row)

    def migrate(self):
        with self.connect() as db:
            if os.getenv('SERVICE_SCHEMA_MANAGED')!='true':
                db.execute(Path('/app/schema.sql').read_text())
                db.execute('ALTER TABLE records DROP CONSTRAINT IF EXISTS records_tenant_kind_body_key')
            ontology=yaml.safe_load(Path('/app/schemas/ontology/v1.yaml').read_text())
            db.execute('INSERT INTO ontology_versions(version,document,actor) VALUES(1,%s,%s) ON CONFLICT DO NOTHING',(Jsonb(ontology),'bootstrap'))
            # Durable work survives a crash. A partial job is replayed idempotently.
            db.execute("UPDATE jobs SET status='pending' WHERE status='running' AND NOT cancelled")
        self.sync_ontology()

    def get(self,ref,db=None):
        try:
            key=str(uuid.UUID(str(ref['id'])))
            version=int(ref['version'])
        except (KeyError,ValueError,TypeError) as exc:
            raise Invalid('Exact version reference required') from exc
        if db is None:
            with self.connect() as conn:
                return self.get(ref,conn)
        row=db.execute('SELECT * FROM records WHERE id=%s AND version=%s AND tenant=%s',(key,version,self.tenant)).fetchone()
        if not row:
            raise Invalid('Referenced version unavailable')
        return row

    def list(self,kind=None,limit=50,offset=0):
        with self.connect() as db:
            rows=db.execute('SELECT DISTINCT ON (id) * FROM records WHERE tenant=%s AND (%s::text IS NULL OR kind=%s) ORDER BY id,version DESC LIMIT %s OFFSET %s',(self.tenant,kind,kind,min(max(limit,1),100),max(offset,0))).fetchall()
            for row in rows:
                review=db.execute("SELECT body->>'decision' AS decision FROM records WHERE kind='Review' AND body->'reviewed'->>'id'=%s AND (body->'reviewed'->>'version')::int=%s ORDER BY created_at DESC LIMIT 1",(str(row['id']),row['version'])).fetchone()
                row['review_state']=review['decision'] if review else 'draft'
            return rows

    def audit(self,db,actor,operation,ref):
        db.execute('INSERT INTO audit(actor,operation,id,version,correlation) VALUES(%s,%s,%s,%s,%s)',(actor,operation,ref['id'],ref['version'],str(uuid.uuid4())))

    def write(self,kind,body,actor,key,previous=None):
        if kind not in KINDS or not key or len(key)>200:
            raise Invalid('Invalid object kind or idempotency key')
        if kind=='Review':
            raise Invalid('Use the human review operation')
        missing=[p for p in REQUIRED.get(kind,[]) if p not in body]
        if missing:
            raise Invalid('Required fields: '+', '.join(missing))
        digest=hashlib.sha256(canonical({'kind':kind,'body':body,'previous':previous})).hexdigest()
        with self.connect() as db:
            db.execute('SELECT pg_advisory_xact_lock(hashtextextended(%s,0))',(key,))
            retry=db.execute('SELECT * FROM requests WHERE key=%s',(key,)).fetchone()
            if retry:
                if retry['digest']!=digest:
                    raise Invalid('Idempotency key already used for different content')
                return self.get(retry,db)
            for field in ('evidence','inputs','findings','artifacts','outputs','definitions'):
                if field=='outputs' and isinstance(body.get(field),dict):
                    raise Invalid('outputs must be a list of exact registered object references; put calculation values in parameters/results')
                for ref in body.get(field,[]):
                    target=self.get(ref,db)
                    if field=='evidence':
                        if target['kind'] not in ('Source','Evidence') or not ref.get('locator'):
                            raise Invalid('Evidence must reference an exact source/evidence version and locator')
                        if target['kind']=='Source':
                            chunk=db.execute('SELECT id FROM chunks WHERE source_id=%s AND source_version=%s AND locator=%s',(ref['id'],ref['version'],ref['locator'])).fetchone()
                            if not chunk:
                                raise Invalid('Evidence locator has no successful extraction')
            if body.get('run'):
                if self.get(body['run'],db)['kind']!='AnalysisRun':
                    raise Invalid('Producing run must be an AnalysisRun')
            if kind=='ReviewerValidation':
                if not actor.startswith('agent:reviewer'):raise Invalid('Independent reviewer identity required')
                if self.get(body['reviewed'],db)['kind']!='Deliverable':raise Invalid('Reviewer validation requires exact deliverable')
                if not body['report'].strip() or not body['reviewer_run']:raise Invalid('Reviewer report and execution identity required')
                if not isinstance(body['checks'],dict) or not body['checks']:raise Invalid('Reviewer must record explicit checks')
            if kind in ('Source','Artifact'):
                data=self.storage.get(body['hash'])
                if len(data)!=body['size']:
                    raise Invalid('Artifact size mismatch')
            if kind=='Assertion':
                ontology=db.execute('SELECT document FROM ontology_versions ORDER BY version DESC LIMIT 1').fetchone()['document']
                if body['type'] not in ontology['types'] or not 0<=body['confidence']<=1:
                    raise Invalid('Invalid assertion type/confidence')
            if kind=='Relation':
                ontology=db.execute('SELECT document FROM ontology_versions ORDER BY version DESC LIMIT 1').fetchone()['document']
                relation=ontology['relations'].get(body['relation'])
                if not relation:raise Invalid('Relation requires approved ontology semantics')
                source=self.get(body['from'],db);target=self.get(body['to'],db)
                source_type=source['body']['type'] if source['kind']=='Assertion' else source['kind']
                target_type=target['body']['type'] if target['kind']=='Assertion' else target['kind']
                if source_type not in relation['from'] or target_type not in relation['to']:raise Invalid('Relation endpoint types are not allowed')
                if relation.get('evidence') and not body['evidence']:raise Invalid('Relation needs exact source evidence')
                for prop in relation.get('required_properties',[]):
                    if prop not in body.get('properties',{}):raise Invalid('Relation missing required property: '+prop)
            if previous:
                old=self.get(previous,db)
                db.execute('SELECT pg_advisory_xact_lock(hashtextextended(%s,0))',(str(old['id']),))
                latest=db.execute('SELECT max(version) AS version FROM records WHERE id=%s',(old['id'],)).fetchone()['version']
                if old['kind']!=kind or latest!=old['version']:
                    raise Invalid('Revision conflict; refresh the exact current version')
                ident,version=old['id'],latest+1
            else:
                ident,version=str(uuid.uuid4()),1
            row=db.execute('INSERT INTO records(id,version,tenant,kind,body,actor) VALUES(%s,%s,%s,%s,%s,%s) RETURNING *',(ident,version,self.tenant,kind,Jsonb(body),actor)).fetchone()
            db.execute('INSERT INTO requests(key,digest,id,version) VALUES(%s,%s,%s,%s)',(key,digest,ident,version))
            db.execute('INSERT INTO outbox(id,version) VALUES(%s,%s)',(ident,version))
            self.audit(db,actor,'create-version',row)
            return row

    def upload(self,data,filename,media,metadata,actor,key):
        if not data or len(data)>20*1024*1024:
            raise Invalid('Source must be between 1 byte and 20 MiB')
        from datetime import datetime,timezone
        digest=self.storage.put(data)
        body={**metadata,'filename':Path(filename).name,'hash':digest,'size':len(data),'media_type':media,
              'retrieved_at':metadata.get('retrieved_at') or datetime.now(timezone.utc).isoformat(),
              'access_limitations':metadata.get('access_limitations','unknown')}
        # Stable ingestion retries do not change their retrieval timestamp.
        with self.connect() as db:
            retry=db.execute('SELECT * FROM requests WHERE key=%s',(key,)).fetchone()
            if retry:
                prior=self.get(retry,db)
                if prior['body']['hash']!=digest or prior['body']['filename']!=body['filename']:
                    raise Invalid('Upload retry differs')
                return prior
        return self.write('Source',body,actor,key)

    def ingest(self,ref):
        row=self.get(ref)
        if row['kind']!='Source':
            raise Invalid('Ingestion requires a source')
        with self.connect() as db:
            job=db.execute('INSERT INTO jobs(id,source_id,version) VALUES(%s,%s,%s) ON CONFLICT(source_id,version) DO UPDATE SET updated_at=now() RETURNING *',(str(uuid.uuid4()),row['id'],row['version'])).fetchone()
            return job

    def extract(self,body):
        data=self.storage.get(body['hash']); name=body['filename'].lower()
        if name.endswith('.pdf'):
            from pypdf import PdfReader
            result=[(f'page:{i+1}',page.extract_text() or '') for i,page in enumerate(PdfReader(io.BytesIO(data)).pages)]
            if not any(text.strip() for _,text in result):
                raise Invalid('Scanned PDF requires OCR; no indexed text')
            return result,'pypdf-6.1.1'
        if name.endswith('.docx'):
            from docx import Document
            doc=Document(io.BytesIO(data))
            return [(f'paragraph:{i+1}',p.text) for i,p in enumerate(doc.paragraphs) if p.text.strip()]+[(f'table:{i+1}', '\n'.join(','.join(c.text for c in r.cells) for r in t.rows)) for i,t in enumerate(doc.tables)],'python-docx-1.2.0'
        if name.endswith('.xlsx'):
            from openpyxl import load_workbook
            book=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
            return [(f'sheet:{sheet.title}:row:{i+1}',','.join('' if c is None else str(c) for c in row)) for sheet in book for i,row in enumerate(sheet.iter_rows(values_only=True))],'openpyxl-3.1.5'
        text=data.decode('utf-8-sig')
        if name.endswith('.csv'):
            return [(f'row:{i+1}',','.join(row)) for i,row in enumerate(csv.reader(io.StringIO(text)))],'csv-v1'
        if name.endswith(('.html','.htm')) or body['media_type']=='text/html':
            from bs4 import BeautifulSoup
            soup=BeautifulSoup(text,'html.parser')
            for bad in soup(['script','style','noscript']): bad.decompose()
            text=soup.get_text('\n',strip=True)
        elif not name.endswith(('.txt','.md','.markdown')):
            raise Invalid('Unsupported extraction format')
        return [(f'section:{i//1800+1}',text[i:i+1800]) for i in range(0,len(text),1800)],'text-v1'

    def embedding_space(self):
        # Provider changes invalidate vectors even when a model alias is reused.
        config={k:os.getenv(k,'') for k in ('EMBEDDING_PROVIDER','EMBEDDING_BASE_URL','EMBEDDING_MODEL','EMBEDDING_DIMENSIONS')}
        config['extraction']='chunk-v1-1800'
        return hashlib.sha256(canonical(config)).hexdigest()

    def embedding(self,texts):
        model=os.getenv('EMBEDDING_MODEL'); base=os.getenv('EMBEDDING_BASE_URL'); dims=os.getenv('EMBEDDING_DIMENSIONS')
        if not model or not base or not dims:
            raise Invalid('Dense retrieval requires configured embedding endpoint, model and dimensions')
        key=os.getenv(os.getenv('EMBEDDING_API_KEY_ENV','DD_EMBEDDING_API_KEY'),'')
        if not key:raise Invalid('Embedding provider requires its configured named key')
        try:
            response=httpx.post(base.rstrip('/')+'/embeddings',json={'model':model,'input':texts},headers={'Authorization':'Bearer '+key},timeout=60,trust_env=False)
            response.raise_for_status()
            values=[r['embedding'] for r in sorted(response.json()['data'],key=lambda x:x['index'])]
            if len(values)!=len(texts) or any(len(v)!=int(dims) for v in values):
                raise Invalid('Embedding dimension/count mismatch')
            return values
        except (httpx.HTTPError,KeyError,ValueError) as exc:
            raise Invalid('Embedding provider failed; no dense index success claimed') from exc

    def process_job(self,job):
        try:
            row=self.get({'id':job['source_id'],'version':job['version']})
            parts,extractor=self.extract(row['body'])
            for i,(locator,text) in enumerate(parts):
                if not text.strip(): continue
                with self.connect() as db:
                    if db.execute('SELECT cancelled FROM jobs WHERE id=%s',(job['id'],)).fetchone()['cancelled']:
                        db.execute("UPDATE jobs SET status='cancelled',updated_at=now() WHERE id=%s",(job['id'],))
                        return
                    db.execute('INSERT INTO chunks(source_id,source_version,locator,text,hash,extractor,configuration) VALUES(%s,%s,%s,%s,%s,%s,%s) ON CONFLICT DO NOTHING',(row['id'],row['version'],locator,text,hashlib.sha256(text.encode()).hexdigest(),extractor,Jsonb({'chunk_size':1800,'version':1})))
                    db.execute('UPDATE jobs SET progress=%s,updated_at=now() WHERE id=%s',(round(100*(i+1)/max(len(parts),1)),job['id']))
            if not parts:
                raise Invalid('Extraction produced no content')
            with self.connect() as db:
                db.execute("UPDATE jobs SET status='extracted',progress=100,updated_at=now() WHERE id=%s",(job['id'],))
            if os.getenv('EMBEDDING_MODEL'):
                self.reindex(row['id'],row['version'])
                with self.connect() as db:
                    db.execute("UPDATE jobs SET status='indexed',updated_at=now() WHERE id=%s",(job['id'],))
        except Exception as exc:
            with self.connect() as db:
                db.execute("UPDATE jobs SET status='failed',error=%s,updated_at=now() WHERE id=%s",(str(exc) if isinstance(exc,Invalid) else 'Extraction/indexing failed; see administrator check',job['id']))

    def reindex(self,source=None,version=None):
        with self.connect() as db:
            rows=db.execute('SELECT * FROM chunks WHERE (%s::uuid IS NULL OR source_id=%s) AND (%s::int IS NULL OR source_version=%s)',(source,source,version,version)).fetchall()
        for start in range(0,len(rows),32):
            batch=rows[start:start+32]
            values=self.embedding([row['text'] for row in batch])
            with self.connect() as db:
                for row,value in zip(batch,values):
                    db.execute('UPDATE chunks SET embedding=%s::vector,embedding_model=%s WHERE id=%s',(str(value),self.embedding_space(),row['id']))
        return {'indexed':len(rows),'model':os.environ['EMBEDDING_MODEL'],'space':self.embedding_space()}

    def search(self,query,limit=20,dense=True):
        if not query.strip() or len(query)>2000:
            raise Invalid('Query must contain 1–2000 characters')
        limit=min(max(limit,1),50)
        with self.connect() as db:
            bm=db.execute('SELECT *,paradedb.score(id) AS score FROM chunks WHERE text @@@ %s ORDER BY score DESC LIMIT %s',(query,limit*2)).fetchall()
            vector=[]
            if dense:
                value=self.embedding([query])[0]
                vector=db.execute('SELECT *,embedding <=> %s::vector AS distance FROM chunks WHERE embedding_model=%s AND vector_dims(embedding)=%s ORDER BY distance LIMIT %s',(str(value),self.embedding_space(),int(os.environ['EMBEDDING_DIMENSIONS']),limit*2)).fetchall()
        scores={};hits={}
        for ranking in (bm,vector):
            for rank,row in enumerate(ranking):
                scores[row['id']]=scores.get(row['id'],0)+1/(60+rank+1)
                row.pop('embedding',None); hits[row['id']]=row
        ranked=[dict(hits[i],rrf=scores[i]) for i in sorted(scores,key=scores.get,reverse=True)[:limit]]
        expansion=[]
        for row in ranked[:3]:
            try:expansion.append(self.graph(str(row['source_id'])+':'+str(row['source_version']),depth=1,limit=10))
            except Exception:expansion.append({'unavailable':True,'bounded':True})
        return {'mode':'hybrid' if dense else 'bm25','hits':ranked,'graph_expansion':expansion,'ranking_counts':{'bm25':len(bm),'dense':len(vector)},'embedding_space':self.embedding_space() if dense else None}

    def web_search(self,query):
        if os.getenv('PUBLIC_RESEARCH_AUTHORIZED')!='true':raise Invalid('Public research requires engagement authorization')
        if not query.strip() or len(query)>2000:raise Invalid('Search query must contain 1–2000 characters')
        provider=os.getenv('WEB_SEARCH_PROVIDER')
        if provider=='bing_rss':
            from urllib.parse import urlencode
            import xml.etree.ElementTree as ET
            data,_,_=fetch_public('https://www.bing.com/search?'+urlencode({'format':'rss','q':query}),limit=1024*1024)
            try:
                root=ET.fromstring(data)
                results=[{'title':item.findtext('title') or '', 'url':item.findtext('link') or '', 'description':item.findtext('description') or ''} for item in root.findall('./channel/item')][:5]
            except ET.ParseError as exc:raise Invalid('Search provider returned invalid RSS') from exc
        elif provider=='brave':
            key=os.getenv(os.getenv('WEB_SEARCH_API_KEY_ENV','DD_SEARCH_API_KEY'),'')
            try:
                response=httpx.get('https://api.search.brave.com/res/v1/web/search',params={'q':query,'count':5},headers={'X-Subscription-Token':key},timeout=30,trust_env=False)
                response.raise_for_status()
                results=[{'title':r['title'],'url':r['url'],'description':r.get('description','')} for r in response.json().get('web',{}).get('results',[])][:5]
            except (httpx.HTTPError,KeyError,ValueError) as exc:raise Invalid('Search provider check failed') from exc
        else:raise Invalid('Configure supported Bing RSS or Brave external search')
        return {'provider':provider,'discovery_only':True,'results':results}

    def validation(self,ref):
        obj=self.get(ref); errors=[]
        if obj['kind']!='Deliverable':
            errors.append('Publication requires a Deliverable')
            return {'valid':False,'errors':errors}
        body=obj['body']
        if not body['findings'] or not body['artifacts']:
            errors.append('Deliverable needs findings and output bytes')
        try:
            run=self.get(body['run'])
            manifest=run['body']
            if manifest['execution_status']!='succeeded' or not manifest['code_commit'] or not manifest['environment'] or not manifest['inputs']:
                errors.append('Incomplete successful run manifest')
            if not manifest.get('outputs'):errors.append('Final run manifest must reference registered outputs')
            if not isinstance(manifest.get('environment'),dict) or not str(manifest['environment'].get('image_digest','')).startswith('sha256:'):errors.append('Run needs pinned runtime image digest')
            with self.connect() as db:
                attested=db.execute("SELECT 1 FROM records WHERE kind='RunAttestation' AND actor='administrator:native-api' AND body->'analysis'->>'id'=%s AND (body->'analysis'->>'version')::int=%s AND body->>'paperclip_run_id'=%s AND body->>'task_id'=%s AND body->>'code_commit'=%s",(str(run['id']),run['version'],manifest['paperclip_run_id'],manifest['task_id'],manifest['code_commit'])).fetchone()
            if not attested:errors.append('Administrator native task/run and Git provenance verification required')
            for input_ref in manifest['inputs']: self.get(input_ref)
            for artifact_ref in body['artifacts']:
                art=self.get(artifact_ref)
                if art['kind']!='Artifact': errors.append('Output reference must be an artifact')
                elif len(self.storage.get(art['body']['hash']))!=art['body']['size']: errors.append('Output checksum/size failed')
            for finding_ref in body['findings']:
                finding=self.get(finding_ref)
                if finding['kind']!='Finding' or not finding['body']['evidence']:
                    errors.append('Finding has no supporting evidence; keep as hypothesis')
                    continue
                for evidence_ref in finding['body']['evidence']:
                    source=self.get(evidence_ref)
                    if source['kind']=='Source': self.storage.get(source['body']['hash'])
                    with self.connect() as db:
                        if source['kind']=='Source' and not db.execute('SELECT id FROM chunks WHERE source_id=%s AND source_version=%s AND locator=%s',(evidence_ref['id'],evidence_ref['version'],evidence_ref.get('locator'))).fetchone():
                            errors.append('Evidence locator unavailable')
                for metric in finding['body'].get('metrics',[]):
                    if not metric.get('definition') or self.get(metric['definition'])['kind']!='Definition':
                        errors.append('Material metric lacks registered definition')
        except Invalid as exc:
            errors.append(str(exc))
        with self.connect() as db:
            pending=db.execute("SELECT count(*) AS n FROM outbox WHERE state!='applied'").fetchone()['n']
            review=db.execute("SELECT body FROM records WHERE kind='ReviewerValidation' AND body->'reviewed'->>'id'=%s AND (body->'reviewed'->>'version')::int=%s ORDER BY created_at DESC LIMIT 1",(str(obj['id']),obj['version'])).fetchone()
        if not review:errors.append('Independent reviewer validation of exact deliverable version required')
        else:
            def failed_check(value):
                if isinstance(value,dict):return value.get('passed') is False or any(failed_check(v) for v in value.values())
                if isinstance(value,list):return any(failed_check(v) for v in value)
                return False
            if review['body'].get('approved') is False or review['body']['checks'].get('blocking_errors') or review['body']['checks'].get('approved') is False or failed_check(review['body']['checks']):errors.append('Reviewer reported blocking validation errors')
        if pending: errors.append('Graph projection has pending or failed changes')
        return {'valid':not errors,'errors':errors,'reviewed':ref}

    def review(self,ref,decision,actor,role,report):
        if role not in ('owner','reviewer') or actor.startswith('agent:'):
            raise Invalid('Human owner/reviewer permission required')
        if decision not in ('accept','reject') or not report.strip():
            raise Invalid('Decision and review rationale required')
        validation=self.validation(ref)
        if decision=='accept' and not validation['valid']:
            raise Invalid('; '.join(validation['errors']))
        with self.connect() as db:
            row=self.get(ref,db)
            ident=str(uuid.uuid4())
            body={'reviewed':ref,'decision':decision,'validation':validation,'report':report}
            review=db.execute("INSERT INTO records(id,version,tenant,kind,status,body,actor) VALUES(%s,1,%s,'Review',%s,%s,%s) RETURNING *",(ident,self.tenant,decision,Jsonb(body),actor)).fetchone()
            # Acceptance is an immutable Review; the reviewed revision stays byte-for-byte unchanged.
            self.audit(db,actor,'human-'+decision,row)
            db.execute('INSERT INTO outbox(id,version) VALUES(%s,1)',(ident,))
            return review

    def graph_driver(self):
        return GraphDatabase.driver(os.getenv('GRAPH_URL','bolt://graph:7687'),auth=None)

    def ontology(self,document,actor,approve=False,role='agent'):
        if not isinstance(document,dict) or not isinstance(document.get('types'),list) or not isinstance(document.get('relations'),dict):raise Invalid('Ontology requires types and relations')
        for name,rule in document['relations'].items():
            if not name.isidentifier() or not name.isupper() or not isinstance(rule,dict):raise Invalid('Invalid relation schema')
            if any(t not in document['types'] for t in rule.get('from',[])+rule.get('to',[])) or not rule.get('from') or not rule.get('to'):raise Invalid('Invalid ontology endpoint types')
        if not approve:
            proposal=self.write('OntologyProposal',{'document':document,'rationale':'Agent proposed schema; unavailable for ingestion until human owner approval'},actor,'ontology-proposal:'+hashlib.sha256(canonical(document)).hexdigest())
            return {'state':'proposal','proposal':proposal,'requires':'human owner approval before use'}
        if role!='owner' or actor.startswith('agent:'):raise Invalid('Ontology approval requires human owner')
        with self.connect() as db:
            db.execute("SELECT pg_advisory_xact_lock(hashtextextended('ontology-approval',0))")
            current=db.execute('SELECT version,document FROM ontology_versions ORDER BY version DESC LIMIT 1').fetchone()
            if {k:v for k,v in current['document'].items() if k!='version'}=={k:v for k,v in document.items() if k!='version'}:
                return {'state':'approved','version':current['version'],'document':current['document']}
            version=db.execute('SELECT coalesce(max(version),0)+1 AS v FROM ontology_versions').fetchone()['v']
            document={**document,'version':version}
            db.execute('INSERT INTO ontology_versions(version,document,actor) VALUES(%s,%s,%s)',(version,Jsonb(document),actor))
            db.execute('INSERT INTO audit(actor,operation,version,correlation) VALUES(%s,%s,%s,%s)',(actor,'approve-ontology',version,str(uuid.uuid4())))
        self.sync_ontology()
        return {'state':'approved','version':version,'document':document}

    def sync_ontology(self):
        """Canonical approved schemas are replayable into a dedicated tenant Git volume."""
        import subprocess
        root=Path(os.getenv('ONTOLOGY_REPO','/ontology'));root.mkdir(parents=True,exist_ok=True)
        def git(*args):
            return subprocess.run(['git','-C',str(root),*args],check=True,capture_output=True,text=True)
        if not (root/'.git').exists():
            git('init');git('config','user.name','Governed DD schema service');git('config','user.email','schema@example.invalid')
        with self.connect() as db:versions=db.execute('SELECT * FROM ontology_versions ORDER BY version').fetchall()
        for row in versions:
            file=root/('v'+str(row['version'])+'.yaml')
            content=yaml.safe_dump(row['document'],sort_keys=True)
            if file.exists():
                if file.read_text()!=content:raise Invalid('Ontology Git export differs from canonical approved schema')
                continue
            file.write_text(content);git('add',file.name);git('commit','-m','Approved ontology v'+str(row['version'])+' by '+row['actor'])

    def project_one(self,event):
        row=self.get(event)
        with self.graph_driver() as driver, driver.session() as graph:
            key=f"{row['id']}:{row['version']}"
            graph.run('MERGE (n:Governed {key:$key}) SET n.id=$id,n.version=$version,n.kind=$kind,n.title=$title',key=key,id=str(row['id']),version=row['version'],kind=row['kind'],title=row['body'].get('claim',row['body'].get('name',row['body'].get('title',row['body'].get('filename',row['kind']))))).consume()
            for field,relation in [('evidence','DERIVED_FROM'),('inputs','USES'),('findings','ABOUT'),('artifacts','DERIVED_FROM')]:
                for ref in row['body'].get(field,[]):
                    other=self.get(ref)
                    target=f"{other['id']}:{other['version']}"
                    graph.run(f'MERGE (a:Governed {{key:$a}}) MERGE (b:Governed {{key:$b}}) MERGE (a)-[r:{relation}]->(b) SET r.locator=$locator',a=key,b=target,locator=ref.get('locator','')).consume()
            if row['body'].get('run'):
                ref=row['body']['run']; self.get(ref)
                graph.run('MATCH (a:Governed {key:$a}) MERGE (b:Governed {key:$b}) MERGE (a)-[:GENERATED_BY]->(b)',a=key,b=f"{ref['id']}:{ref['version']}").consume()
            if row['kind']=='Relation':
                body=row['body'];relation=body['relation']
                if not relation.isidentifier() or not relation.isupper():raise Invalid('Invalid projected relation')
                a=body['from'];b=body['to']
                graph.run(f'MERGE (a:Governed {{key:$a}}) MERGE (b:Governed {{key:$b}}) MERGE (a)-[r:{relation} {{record:$record}}]->(b) SET r.evidence=$evidence',a=f"{a['id']}:{a['version']}",b=f"{b['id']}:{b['version']}",record=key,evidence=json.dumps(body['evidence'])).consume()

    def project(self,rebuild=False):
        if rebuild:
            with self.connect() as db: db.execute("UPDATE outbox SET state='pending',error=NULL")
        with self.connect() as db:
            events=db.execute("SELECT * FROM outbox WHERE state IN ('pending','failed') ORDER BY sequence LIMIT 100").fetchall()
        for event in events:
            try:
                self.project_one(event)
                with self.connect() as db: db.execute("UPDATE outbox SET state='applied',attempts=attempts+1,error=NULL WHERE sequence=%s",(event['sequence'],))
            except Exception:
                with self.connect() as db: db.execute("UPDATE outbox SET state='failed',attempts=attempts+1,error='Graph projection failed' WHERE sequence=%s",(event['sequence'],))
        return {'processed':len(events)}

    def graph(self,key,depth=1,limit=40,relation=None):
        depth=min(max(depth,1),3); limit=min(max(limit,1),100)
        with self.connect() as db:allowed=set(db.execute('SELECT document FROM ontology_versions ORDER BY version DESC LIMIT 1').fetchone()['document']['relations'])
        if relation and relation not in allowed: raise Invalid('Unsupported relation filter')
        rel=':'+relation if relation else ''
        with self.graph_driver() as driver,driver.session() as graph:
            result=graph.run(f'MATCH p=(n:Governed {{key:$key}})-[{rel}*1..{depth}]-(m) RETURN p LIMIT $limit',key=key,limit=limit)
            nodes={}; edges={}
            for record in result:
                for node in record['p'].nodes: nodes[node['key']]=dict(node)
                for edge in record['p'].relationships:
                    edge_key=(edge.start_node['key'],edge.end_node['key'],edge.type)
                    edges[edge_key]={'from':edge_key[0],'to':edge_key[1],'type':edge.type,**dict(edge)}
            return {'nodes':list(nodes.values()),'edges':list(edges.values()),'bounded':True}
